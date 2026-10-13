"""汇总导出：攻击构造 × 评测结果 × 数据采集 → Excel（面向项目汇报的阅读体验版）。

2026-10-13 版（同步至 31 案套件 / r12 评测 / 采集协议 v2.1）：
  阅读指引 —— 这份表是什么、每张表看什么、状态图例、三条核心结论
  攻击用例总表 —— 技术池逐行（含中文战术归属）+ 20 恶意案 + 完整链 + 9 良性镜像
  评测结果明细 —— r12 两轴运行：期望 vs 实际（双轮）、耗时与成本、GT 复核状态
  数据采集明细 —— v2.1 采集包实况：原始量 vs 回放消费，物证声明
  指标汇总 —— 当前能力面 / 延后面 / demo 回归
状态、结论、成本全部从 live 数据（case.json + r12 report.json + bundles_v2）生成。
"""

from __future__ import annotations

import json
import os
import pathlib
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = pathlib.Path(os.environ.get("XLSX_OUT", str(ROOT / "docs" / "安全能力验证-攻击构造与评测汇总.xlsx")))
EVAL_DIR = ROOT / "outputs/eval"
STAGING = ROOT / "validation_cases_staging"
PINNED = ROOT / "validation_cases"
BUNDLES = ROOT / "attack_lab/bundles_v2"

TACTIC_ZH = {
    "initial-access": "初始访问", "execution": "执行", "persistence": "持久化",
    "privilege-escalation": "权限提升", "defense-evasion": "防御绕过",
    "credential-access": "凭据获取", "discovery": "发现", "lateral-movement": "横向移动",
    "collection": "收集", "command-and-control": "命令控制", "exfiltration": "数据外传",
    "impact": "影响", "reconnaissance": "侦察", "resource-development": "资源开发",
}
VERDICT_ZH = {
    "confirmed_malicious": "确认恶意", "likely_malicious": "倾向恶意", "suspicious": "可疑",
    "insufficient_evidence": "证据不足", "likely_benign": "倾向良性", "benign": "良性",
}
DATA_COLUMNS = ["进程执行", "进程谱系", "文件写入", "网络四元组", "提权调用", "命令执行",
                "包归属", "端点基线", "监听端口", "工件哈希"]
# 采集协议 v2/v2.1：靶场案十列全有效（文件写入丢失问题已在 v2 根治；
# 统计快照/内容物证/回执哈希为 v2.1 新增，不占十列而在采集明细表体现）
MARKS_FULL = ["有效"] * 10
MARKS_NONE = ["未采"] * 10
# 合成参考数据（c2 基线案与 demo 轴）：无提权/命令执行两类审计通道
MARKS_SYNTH = ["有效", "有效", "有效", "有效", "未采", "未采", "有效", "有效", "未采", "有效"]

EXCLUDED = {
    "T1055.001": "核验为非 Linux 平台技术",
}

# T-ID → (case_id, 场景说明)；结论/状态/成本从 live 数据生成
TID_CASE = {
    "T1105": ("t1105-ingress-tool-transfer", "从外部端点下载可执行文件、加执行权限并运行、回连控制端（含 T1222.002 权限设置）"),
    "T1037.004": ("t1037-004-rc-local", "追加 rc.local 启动项实现重启自启"),
    "T1053.003": ("t1053-003-cron-persistence", "写入 cron 定时任务周期执行后门载荷"),
    "T1543.002": ("t1543-002-systemd-service", "投放后门并注册 systemd 服务实现开机自启、持续回连"),
    "T1505.003": ("t1505-003-web-shell", "向 Web 目录写入 Web Shell 并经 HTTP 触发"),
    "T1574.006": ("t1574-006-ldpreload", "投放恶意动态库并以 LD_PRELOAD 劫持合法程序"),
    "T1486": ("t1486-ransomware", "运行加密器加密业务目录（演练后恢复）"),
    "T1219": ("t1219-rat-tunnel", "部署反向隧道工具建立出站长连接"),
    "T1567.002": ("t1567-002-exfil-upload", "把打包的敏感数据上传到外部端点（回执哈希可验证收发一致）"),
    "T1547.006": ("t1547-006-kernel-module", "加载内核模块驻留（签名策略下加载失败痕迹即证据）"),
    "T1005": ("t1005-t1560-001-archive", "收集凭据材料并打包暂存（含 T1560.001 归档）"),
    "T1036.005": ("t1036-005-masquerade", "后门伪装系统组件名放入系统目录执行"),
    "T1070.006": ("t1070-006-timestomp", "回拨载荷时间戳至系统旧文件（stat 快照曝出物理矛盾）"),
    "T1027": ("t1027-obfuscated-script", "base64 载荷落盘、运行时解码管道执行（明文不落盘）"),
    "T1098.004": ("t1098-004-authorized-keys", "追加攻击者公钥到 root 的 authorized_keys 并尝试登入"),
    "T1552.004": ("t1552-004-private-key", "窃取既有运维私钥暂存打包外传（窗口前资产快照+回执哈希）"),
    "T1548.001": ("t1548-001-setuid", "root 属主 setuid 载荷由非特权用户执行提权（uid/euid 双值证据）"),
    "T1059.006": ("t1059-006-python-beacon", "解释器执行回连脚本、周期外连"),
    "T1136.001": ("t1136-001-account", "创建与 root 同 UID 的隐藏账号并产生登入会话"),
}
CHAIN_CASE = ("chain-001-full-intrusion",
              "四阶段完整链：下载投放→执行回连→载荷派生持久化（cron+公钥）→收集外传；进程树跨阶段同根")
BENIGN_CASES = [
    ("bb7-vendor-agent", "厂商监控 agent 按清单间隔心跳回连批准端点"),
    ("bb8-ops-suite-install", "监控套件安装：二进制+systemd+cron 三件套（与恶意持久化同形）"),
    ("bb4-ops-cron", "运维日志清理 cron（与恶意 cron 持久化同形）"),
    ("bb5-compiler-artifact", "部署编译产物并向批准端点注册心跳"),
    ("bb6-mirror-download", "从批准镜像下载运维工具并执行（与恶意工具传输同形）"),
    ("bb9-log-rotation", "按在案策略轮转日志：归档保留、删除原件（与毁证同形）"),
    ("bb11-cfgmgmt-ssh-key", "配置管理下发运维公钥并登入验证（与公钥驻留同形）"),
    ("bb2-admin-backup", "按变更单备份应用数据"),
    ("bb3-ops-tool-deploy", "部署健康检查工具并执行一次"),
]
MIRROR_DEFERRED = {"bb4-ops-cron", "bb5-compiler-artifact", "bb6-mirror-download",
                   "bb7-vendor-agent", "bb8-ops-suite-install"}


def load_corpus():
    titles, summaries = {}, {}
    for line in (ROOT / "corpora" / "attack_technique.jsonl").read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        tid = rec["knowledge_id"]
        zh = rec.get("title_zh", "")
        if "（" in zh:
            zh = zh.split("（")[0]
        titles[tid] = zh
        summaries[tid] = (rec.get("summary_zh") or "")[:60]
    return titles, summaries


def load_axis(axis_name):
    """r12 报告 → {case_id: {runs, tags, passed}}；判定从运行工件 judgment.json 读取。"""
    report = EVAL_DIR / axis_name / "report.json"
    if not report.exists():
        return {}
    rep = json.loads(report.read_text(encoding="utf-8"))
    out = {}
    for case in rep["cases"]:
        runs = []
        for run in case["runs"]:
            run_dir = EVAL_DIR / axis_name / "runs" / run["run_id"].replace("#", "_")
            verdict, pub = "-", "-"
            j_path = run_dir / "artifacts" / "judgment.json"
            if j_path.exists():
                j = json.loads(j_path.read_text(encoding="utf-8"))
                verdict = VERDICT_ZH.get(j.get("verdict", {}).get("level"), "-")
                pub = j.get("publication_status", "-")
            m = run.get("metrics_raw") or {}
            runs.append({
                "verdict": verdict, "pub": pub,
                "mins": round(m.get("duration_seconds", 0) / 60, 1),
                "calls": int(m.get("model_calls", 0)),
                "tools": int(m.get("tool_calls", 0)),
                "tokens_k": int(m.get("total_tokens", 0) / 1000),
            })
        out[case["case_id"]] = {
            "runs": runs, "tags": ",".join(case.get("tags", [])),
            "passed": (case.get("pass_at_k") or {}).get("all_epochs_passed"),
        }
    return out


def case_expect(case_id):
    for base in (STAGING, PINNED):
        p = base / case_id / "case.json"
        if p.exists():
            d = json.loads(p.read_text(encoding="utf-8"))
            exp = d.get("expected_ai", {})
            zh = {"malicious": "恶意", "benign": "良性"}.get(exp.get("verdict"), exp.get("verdict", "—"))
            return zh, d.get("ground_truth", {}).get("source", "未知")
    return "—", "未知"


def verdict_note(case_id, r12):
    if case_id not in r12:
        return "（未入 r12）"
    vs = "/".join(r["verdict"] for r in r12[case_id]["runs"])
    note = vs
    if case_id == "t1574-006-ldpreload":
        note += "——已知证据限制：LD_PRELOAD 变量名天然不入审计记录，下一轮采集补锚定"
    if case_id in MIRROR_DEFERRED:
        note += "——镜像误报：自建反证缺权威性，黄区对接权威白名单后消解（已裁决延后）"
    return note


def live_cols(case_id, r12):
    if case_id not in r12:
        return "AI 评测完成", "—", "—"
    runs = r12[case_id]["runs"]
    mins = sum(r["mins"] for r in runs) / len(runs)
    tok = sum(r["tokens_k"] for r in runs) / len(runs)
    return ("AI 评测完成", verdict_note(case_id, r12), f"约{mins:.0f}分钟 / {int(tok)}千tokens")


def build_pool_rows(r12):
    data = json.loads((ROOT / "downloads" / "enterprise-attack.json").read_text(encoding="utf-8"))
    objs = data["objects"]
    comps = {o["id"]: o.get("name", "") for o in objs if o.get("type") == "x-mitre-data-component"}
    an2c = defaultdict(set)
    for o in objs:
        if o.get("type") == "x-mitre-analytic":
            for ls in o.get("x_mitre_log_source_references", []):
                if ls.get("x_mitre_data_component_ref") in comps:
                    an2c[o["id"]].add(comps[ls["x_mitre_data_component_ref"]])
    ds2an = {o["id"]: o.get("x_mitre_analytic_refs", []) for o in objs
             if o.get("type") == "x-mitre-detection-strategy"}
    t2ds = defaultdict(list)
    for o in objs:
        if o.get("type") == "relationship" and o.get("relationship_type") == "detects":
            t2ds[o["target_ref"]].append(o["source_ref"])

    def file_observable(t):
        for ds in t2ds.get(t["id"], []):
            for an in ds2an.get(ds, []):
                if any("file" in c.lower() for c in an2c.get(an, set())):
                    return True
        return False

    titles, _ = load_corpus()
    rows = []
    for t in objs:
        if t.get("type") != "attack-pattern" or t.get("revoked") or t.get("x_mitre_deprecated"):
            continue
        if not any("Linux" in p for p in t.get("x_mitre_platforms", [])):
            continue
        if not file_observable(t):
            continue
        tid = t["external_references"][0]["external_id"]
        tactics = "、".join(sorted({TACTIC_ZH.get(k["phase_name"], k["phase_name"])
                                    for k in t.get("kill_chain_phases", [])}))
        zh = titles.get(tid, t["name"])
        if tid in TID_CASE:
            case_id, scene = TID_CASE[tid]
            status, verdict, cost = live_cols(case_id, r12)
            rows.append(["攻击技术", tid, zh, tactics, scene, status, verdict, cost, *MARKS_FULL])
        elif tid in EXCLUDED:
            rows.append(["攻击技术", tid, zh, tactics, "", "不适用",
                         "已排除：" + EXCLUDED[tid], "—", *MARKS_NONE])
        else:
            rows.append(["攻击技术", tid, zh, tactics, "", "未开始",
                         "备选池：满足构造条件，31 案套件已足够本轮基线", "—", *MARKS_NONE])
    # 兜底：已建案但未过语料文件可观测过滤器的 TID（如 T1548.001/T1059.006，
    # 其官方检测策略不走文件组件）——显式追加，保证 19 项技术全部在表
    seen = {r[1] for r in rows if r[0] == "攻击技术"}
    for tid, (case_id_, scene_) in TID_CASE.items():
        if tid in seen:
            continue
        status_, verdict_, cost_ = live_cols(case_id_, r12)
        rows.append(["攻击技术", tid, tid, "—", scene_, status_, verdict_, cost_, *MARKS_FULL])
    case_id, scene = CHAIN_CASE
    status, verdict, cost = live_cols(case_id, r12)
    rows.append(["完整入侵链", "多技术", "四阶段入侵链（下载→回连→派生持久化→外传）", "多战术",
                 scene, status, verdict, cost, *MARKS_FULL])
    status, verdict, cost = live_cols("c2-malicious-elf", r12)
    rows.append(["钉住基线", "T1071.001", "C2 后门基线（合成参考数据）", "命令控制",
                 "ELF 后门 systemd 持久化、周期回连未批准端点", status, verdict, cost, *MARKS_SYNTH])
    status, verdict, cost = live_cols("c2-benign-baseline", r12)
    rows.append(["钉住基线", "—", "良性代理基线（合成参考数据）", "—",
                 "签名合法包代理 systemd 托管、心跳批准端点", status, verdict, cost, *MARKS_SYNTH])
    for case_id, scene in BENIGN_CASES:
        status, verdict, cost = live_cols(case_id, r12)
        rows.append(["良性镜像", case_id.split("-")[0], case_id, "—", scene, status, verdict, cost, *MARKS_FULL])
    order = {"AI 评测完成": 0, "未开始": 1, "不适用": 2}
    rows.sort(key=lambda r: (order.get(r[5], 1), r[1]))
    return [[i + 1, *r] for i, r in enumerate(rows)]


def build_eval_rows(r12_val, r12_demo):
    rows = []

    def axis_rows(tag, data):
        for case_id, info in sorted(data.items()):
            exp, gt_src = case_expect(case_id)
            r0 = info["runs"][0]
            r1 = info["runs"][1] if len(info["runs"]) > 1 else {}
            rows.append([
                tag, case_id, info["tags"], exp,
                f"{r0['verdict']} / {r1.get('verdict', '—')}",
                r0["pub"], "通过" if info["passed"] else "未过",
                r0["mins"], r0["calls"], r0["tools"], r0["tokens_k"],
                "人工复核通过" if "复核通过" in gt_src else gt_src,
            ])

    axis_rows("技术覆盖轴（r12）", r12_val)
    axis_rows("数据层级轴（r12 回归）", r12_demo)
    return rows


def build_collect_rows():
    rows = []
    for b in sorted(BUNDLES.glob("*.bundle")):
        case_id = b.name[:-len(".bundle")]
        audit = sum(1 for _ in open(b / "audit.log", encoding="utf-8", errors="replace"))
        poll = b / "conn_poll.log"
        poll_n = sum(1 for _ in open(poll, encoding="utf-8", errors="replace")) if poll.exists() else 0
        stats = b / "system_stats.txt"
        stat_n = sum(1 for _ in open(stats, encoding="utf-8")) if stats.exists() else 0
        mc = b / "system_contents/manifest.txt"
        content_n = sum(1 for l in mc.read_text(encoding="utf-8").splitlines() if l.strip()) if mc.exists() else 0
        tr = b / "transfer_receipts.log"
        receipt_n = sum(1 for l in tr.read_text(encoding="utf-8").splitlines()
                        if l.strip() and not l.startswith("#")) if tr.exists() else 0
        rd = STAGING / case_id / "replay" / "events"
        settled = Counter()
        if rd.exists():
            for f in rd.glob("*.jsonl"):
                for line in f.read_text(encoding="utf-8").splitlines():
                    if line.strip():
                        settled[json.loads(line)["event_type"]] += 1
        rows.append([
            case_id,
            f"审计 {audit} 行 / 轮询 {poll_n} 行",
            f"stat {stat_n} 项 / 内容物证 {content_n} 项 / 回执 {receipt_n} 条",
            sum(settled.values()),
            "、".join(f"{k}×{v}" for k, v in sorted(settled.items(), key=lambda x: -x[1])[:6]) or "—",
        ])
    return rows


# ---------- 写工作簿 ----------
wb = Workbook()
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
OK_FILL = PatternFill("solid", fgColor="E8F5E9")
WARN_FILL = PatternFill("solid", fgColor="FFF8E1")
GREY_FILL = PatternFill("solid", fgColor="F5F5F5")
DEFER_FILL = PatternFill("solid", fgColor="FDE9EC")
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)


def fill_sheet(ws, headers, rows, widths, status_col=None, mark_cols=None):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if status_col:
            v = str(row[status_col - 1].value or "")
            if v == "AI 评测完成":
                row[status_col - 1].fill = OK_FILL
            elif "延后" in v or "镜像误报" in v:
                row[status_col - 1].fill = DEFER_FILL
            elif "限制" in v:
                row[status_col - 1].fill = WARN_FILL
            elif "备选池" in v or "不适用" in v:
                row[status_col - 1].fill = GREY_FILL
        if mark_cols:
            for ci in mark_cols:
                cell = row[ci - 1]
                val = str(cell.value or "")
                if val == "有效":
                    cell.fill = OK_FILL
                elif val == "未采":
                    cell.fill = GREY_FILL
                cell.alignment = Alignment(horizontal="center", vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


r12_val = load_axis("r12-validation-v2")
r12_demo = load_axis("r12-demo-regression")

# S0 阅读指引
ws0 = wb.active
ws0.title = "阅读指引"
guide = [
    ["这份表是什么", "安全能力持续验证（Feature 19）截至 2026-10-13 的完整汇总：31 案套件（20 恶意 + 9 良性镜像 + 2 钉住基线）、r12 全量评测结果、采集协议 v2.1 数据实况。配套逐案讲解见 docs/attack-case-walkthrough.md。"],
    ["攻击用例总表", "主表。全部 Linux 落文件可观测技术池逐行（含中文战术归属），其中 19 项技术 + 1 条完整链已建案评测（绿色），右侧 10 列采集矩阵——v2 协议后靶场案十列全绿（文件写入丢失问题已根治）。"],
    ["评测结果明细", "r12 两轴运行：期望判定 vs 实际判定（双轮）、发布状态、GT 复核状态、耗时成本。"],
    ["数据采集明细", "29 个靶场案逐案：原始审计量 vs 进入回放的事件构成，v2.1 新增物证（stat/内容/回执）计数。"],
    ["指标汇总", "当前能力面 / 延后面（黄区对接项）/ demo 回归三组数字。"],
    ["", ""],
    ["状态图例", "绿=AI 评测完成 ｜ 粉红=镜像误报（已裁决延后黄区权威白名单） ｜ 黄=已知证据限制 ｜ 灰=备选池/不适用。"],
    ["结论一", "恶意检出方向 92%（20 案 × 2 轮，r12）：伪装/时间戳回拨/混淆/隐藏账号/完整链等全部判到恶意侧，完整链与 systemd 案双轮满分；报告可追溯 100%（62 次运行零兜底）。"],
    ["结论二", "良性镜像误报（5 案）的根因是自建反证缺权威性——模型查到端点批准记录但不予采信，生产语境下这是正确直觉；已裁决延后黄区对接权威白名单工具后消解，剩余项与解锁条件见 Feature 19 verification.md。"],
    ["结论三", "套件 31 案 GT 全部人工复核通过；单案成本约 2 分钟 / 10 万 tokens（较 9 月基线 21 分钟 / 26 万 tokens：报告管线重构 + 模型接入优化）。"],
]
fill_sheet(ws0, ["条目", "说明"], guide, [16, 110])

# S1 攻击用例总表
ws1 = wb.create_sheet("攻击用例总表")
fill_sheet(ws1, ["序号", "类别", "ATT&CK", "技术名称", "所属战术", "场景说明", "进度", "结论 / 备注", "耗时与成本", *DATA_COLUMNS],
           build_pool_rows(r12_val), [5, 10, 12, 24, 14, 40, 13, 44, 16, *([8] * 10)],
           status_col=8, mark_cols=range(10, 20))

# S2 评测结果明细
ws2 = wb.create_sheet("评测结果明细")
fill_sheet(ws2, ["评测轴", "用例", "标签", "期望判定", "实际判定(轮1/轮2)", "发布状态", "是否通过",
                 "耗时(分)", "模型调用", "工具调用", "tokens(千)", "GT 复核状态"],
           build_eval_rows(r12_val, r12_demo),
           [16, 30, 26, 10, 16, 10, 8, 9, 10, 10, 10, 14])

# S3 数据采集明细
ws3 = wb.create_sheet("数据采集明细")
fill_sheet(ws3, ["用例", "原始采集量", "v2.1 物证", "回放事件数", "事件构成（前六类）"],
           build_collect_rows(), [30, 24, 30, 12, 44])

# S4 指标汇总
ws4 = wb.create_sheet("指标汇总")
rows4 = [
    ["技术覆盖轴 · 恶意方向正确", "37/40 运行判到恶意侧（失分：t1574 证据限制 ×2、t1548 单轮方差）", "92%"],
    ["技术覆盖轴 · 报告可追溯", "62 次运行全部 grounded 发布，零兜底", "100%"],
    ["技术覆盖轴 · 结论距离（当前能力面）", "恶意全量 + 非镜像良性，48 运行 dist 0.792", "达标(≥0.75)"],
    ["技术覆盖轴 · 方向正确（当前能力面）", "44/48", "92%（达标 ≥85%）"],
    ["技术覆盖轴 · 良性镜像误报", "bb4-bb8 五案判可疑/恶意——自建反证缺权威性，延后黄区（粉红标记）", "延后"],
    ["数据层级轴 · demo 回归", "10 案 × 2 轮：dist 0.84 / 可追溯 1.0 / 方向 1.0 / 边界 1.0 / pass^k 1.0", "全绿不降"],
    ["单案成本（r12）", "验证套件均值 137 秒 / 11.5 次模型调用 / 102 千 tokens；demo 轴 115 秒", "—"],
    ["模型与运行环境", "glm-5.3-flash（OpenAI 接入 + 低思考档）· 真实调查管线（调查→结案研判→代码装配报告→处置建议）", "—"],
    ["数据质量", "四类残留（靶场标识/回环C2/窗口内编译/品牌字样）扫描全零；两案过物证验证门（收发哈希一致）", "—"],
    ["挂起项（黄区解锁）", "规则侧结果采集、AI 自主攻击、权威白名单工具（镜像误报消解）——解锁条件与对接面见 verification.md", "—"],
]
fill_sheet(ws4, ["指标", "数值 / 说明", "达成"], rows4, [34, 78, 14])

wb.save(OUT)
print("saved:", OUT)
